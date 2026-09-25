"""ADR 0134, F5: the workbook is the report's model with its formulas written out.

What is under test is the claim the workbook makes by existing:
- **Its formulas reproduce the report.** Recomputed in a real spreadsheet engine, LibreOffice,
  never by Python reading its own output back.
- **Changing an input moves the answer where the platform says it goes.** Terminal growth
  lands on the report's own sensitivity cell, and revenue growth on the calculator's strike
  for the same change.
- **What may not leave, does not.** The vendor's close stays out.

The spreadsheet engine is LibreOffice's headless calculator. Where it is missing the recompute
tests skip, except in CI, which installs it and fails them rather than letting them pass by
not running.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import time
import zipfile
from dataclasses import replace
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from typing import Any, NoReturn
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from aer.calc.comps import (
    Audience,
    CompsTable,
    MultipleBasis,
    MultipleResult,
    PeerExclusion,
    PeerRow,
)
from aer.calc.units import DIMENSIONLESS, Quantity
from aer.core.enums import JobStatus
from aer.core.hashing import sha256_hex
from aer.db.models import Artefact, JobStep, Report, Sensitivity
from aer.render.workbook import (
    WHERE_THE_RECORD_ENDS,
    WORKBOOK_MEDIA_TYPE,
    ModelWorkbook,
    build_workbook,
)
from aer.services.assumptions import confirmed_values
from aer.services.preview import basis_for_run, strike
from aer.services.workbook import NoWorkbook, recorded_comps, workbook_model
from aer.storage.local import LocalArtefactStore
from aer.workflow.engine import StepContext
from aer.workflow.workflows.vertical_slice_v1 import _archive_workbook
from tests.assumption_fixtures import a_year, seed_years
from tests.valued_run_fixtures import SHARES, valued_run

pytestmark = pytest.mark.integration

# CI installs the spreadsheet engine and must run these; a laptop without it skips them.
# Read at import, before the suite's hermetic environment strips anything.
_ENGINE_REQUIRED = bool(os.environ.get("CI"))

# A close nobody would type by accident, so a test can prove it never reaches the file.
_CLOSE = "41.37"
_CAPITALISATION = "4550.7"

# LibreOffice computes in binary floating point, the ledger in decimal: agreement is held to
# a part in a billion, which is far inside a cent on any per-share figure.
_TOLERANCE = 1e-9


async def _approved(session: AsyncSession, scene: dict[str, Any]) -> Report:
    report = Report(
        job_id=scene["job"].id,
        request_id=scene["request"].id,
        company_id=scene["company"].id,
        as_of_date=scene["request"].work_order.as_of_date,
        approved_by=scene["user"].id,
        approved_at=datetime(2026, 9, 25, 12, 0, tzinfo=UTC),
        content={"markdown": "# A report"},
        content_hash=sha256_hex("# A report"),
    )
    session.add(report)
    await session.flush()
    return report


@pytest.fixture
async def written(db_session: AsyncSession) -> dict[str, Any]:
    scene = await valued_run(db_session, price=_CLOSE, capitalisation=_CAPITALISATION)
    report = await _approved(db_session, scene)
    model = await workbook_model(db_session, job=scene["job"], report=report)
    assert isinstance(model, ModelWorkbook), model
    return {**scene, "report": report, "model": model, "bytes": build_workbook(model)}


def _recalculated(data: bytes, tmp_path: Path) -> Workbook:
    """The workbook as a spreadsheet engine computes it, cached values and all."""
    soffice = shutil.which("soffice")
    if soffice is None:
        _unavailable("LibreOffice is not installed, so nothing can recompute the workbook.")
    source = tmp_path / "model.xlsx"
    source.write_bytes(data)
    out = tmp_path / "out"
    subprocess.run(  # noqa: S603 -- a fixed binary on a file this test wrote
        [
            soffice,
            f"-env:UserInstallation={(tmp_path / 'profile').as_uri()}",
            "--headless",
            "--calc",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out),
            str(source),
        ],
        check=False,
        capture_output=True,
        timeout=240,
    )
    converted = out / "model.xlsx"
    if not converted.exists():
        _unavailable(
            "LibreOffice could not open the workbook: its spreadsheet module "
            "(libreoffice-calc) is probably not installed."
        )
    return load_workbook(converted, data_only=True)


def _unavailable(reason: str) -> NoReturn:
    if _ENGINE_REQUIRED:
        pytest.fail(reason)
    pytest.skip(reason)


def _row_labelled(sheet: Worksheet, label: str) -> int:
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == label:
            return row
    message = f"no row labelled {label!r}"
    raise AssertionError(message)


def _per_share(book: Workbook) -> tuple[float, float]:
    sheet = book["Model"]
    row = _row_labelled(sheet, "Value per share")
    gordon = sheet.cell(row=row, column=2).value
    exit_multiple = sheet.cell(row=row, column=3).value
    assert isinstance(gordon, float | int), gordon
    assert isinstance(exit_multiple, float | int), exit_multiple
    return float(gordon), float(exit_multiple)


def _close(actual: float, expected: Decimal) -> bool:
    return abs(actual - float(expected)) <= _TOLERANCE * max(1.0, abs(float(expected)))


def _set_inputs(data: bytes, label: str, value: Decimal) -> bytes:
    """The workbook with every blue cell on the row ``label`` changed to ``value``."""
    book = load_workbook(BytesIO(data))
    sheet = book["Model"]
    row = _row_labelled(sheet, label)
    for column in range(2, sheet.max_column + 1):
        cell = sheet.cell(row=row, column=column)
        colour = cell.font.color
        if cell.value is not None and colour is not None and colour.rgb in ("000000FF", "FF0000FF"):
            cell.value = value
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class TestTheWorkbookIsTheReportsModel:
    async def test_recomputed_it_gives_the_reports_own_figures(
        self, written: dict[str, Any], tmp_path: Path
    ) -> None:
        book = _recalculated(written["bytes"], tmp_path)

        gordon, exit_multiple = _per_share(book)
        recorded = written["model"].recorded
        assert _close(gordon, recorded["gordon_per_share"]), (gordon, recorded)
        assert _close(exit_multiple, recorded["exit_multiple_per_share"]), (exit_multiple, recorded)

    async def test_the_discount_rate_is_the_reports_own(
        self, written: dict[str, Any], tmp_path: Path
    ) -> None:
        """Live from its parts — the cost of equity, the cost of debt, the market weights —
        and still the rate the report discounted at."""
        book = _recalculated(written["bytes"], tmp_path)
        sheet = book["Model"]
        wacc = sheet.cell(row=_row_labelled(sheet, "Discount rate (WACC)"), column=2).value

        basis = await basis_for_run(written["session"], job=written["job"])
        assert basis is not None
        values = await confirmed_values(written["session"], written["request"].id)
        expected = strike(basis, values).capital.wacc.value
        assert isinstance(wacc, float)
        assert _close(wacc, expected)

    async def test_terminal_growth_moves_it_onto_the_reports_own_grid_cell(
        self, written: dict[str, Any], tmp_path: Path
    ) -> None:
        """F5's "done when", on the input the grid actually varies: a growth rate typed into
        the sheet lands the perpetuity value on the cell the report recorded for it, at the
        report's own discount rate."""
        session: AsyncSession = written["session"]
        grid = await session.scalar(
            select(Sensitivity)
            .where(
                Sensitivity.job_id == written["job"].id,
                # Stored with the discount rate as x and the varied assumption as y.
                Sensitivity.y_assumption == "terminal_growth",
            )
            .options(selectinload(Sensitivity.cells))
        )
        assert grid is not None
        basis = await basis_for_run(session, job=written["job"])
        assert basis is not None
        struck = strike(basis, await confirmed_values(session, written["request"].id))
        base_growth = written["model"].terminal_growth.value
        chosen = max(cell.y_value for cell in grid.cells if cell.y_value != base_growth)
        target = next(
            cell.output_value
            for cell in grid.cells
            if cell.y_value == chosen
            and abs(cell.x_value - struck.capital.wacc.value) < Decimal("1e-9")
        )

        changed = _set_inputs(written["bytes"], "Perpetual growth after the forecast", chosen)
        gordon, _ = _per_share(_recalculated(changed, tmp_path))

        assert _close(gordon, target), (gordon, target)

    async def test_revenue_growth_moves_it_where_the_calculator_says(
        self, written: dict[str, Any], tmp_path: Path
    ) -> None:
        """The input the grid does not vary, held to the calculator's strike instead."""
        basis = await basis_for_run(written["session"], job=written["job"])
        assert basis is not None
        values = dict(await confirmed_values(written["session"], written["request"].id))
        held = values["revenue_growth"]
        values["revenue_growth"] = Quantity.of(Decimal("0.08"), held.unit, source=held.source)
        expected = strike(basis, values).result

        changed = _set_inputs(written["bytes"], "Revenue growth", Decimal("0.08"))
        gordon, exit_multiple = _per_share(_recalculated(changed, tmp_path))

        assert _close(gordon, expected.gordon.value_per_share.value)
        assert _close(exit_multiple, expected.exit_multiple.value_per_share.value)


class TestWhatTheWorkbookCarries:
    def _book(self, written: dict[str, Any]) -> Workbook:
        return load_workbook(BytesIO(written["bytes"]))

    async def test_every_computed_cell_is_a_formula(self, written: dict[str, Any]) -> None:
        sheet = self._book(written)["Model"]
        for label in ("Revenue", "Free cash flow to the firm", "Value per share"):
            row = _row_labelled(sheet, label)
            written_cells = [
                sheet.cell(row=row, column=column).value
                for column in range(2, sheet.max_column + 1)
            ]
            formulas = [value for value in written_cells if isinstance(value, str)]
            assert formulas, label
            assert all(value.startswith("=") for value in formulas), label

    async def test_inputs_are_blue_and_unlocked_formulas_black_and_locked(
        self, written: dict[str, Any]
    ) -> None:
        sheet = self._book(written)["Model"]
        rate = sheet.cell(row=_row_labelled(sheet, "Risk-free rate"), column=2)
        wacc = sheet.cell(row=_row_labelled(sheet, "Discount rate (WACC)"), column=2)

        assert rate.font.color is not None
        assert rate.font.color.rgb.endswith("0000FF")
        assert rate.protection.locked is False
        assert wacc.font.color is not None
        assert wacc.font.color.rgb.endswith("000000")
        assert wacc.protection.locked is True
        assert sheet.protection.sheet is True
        assert not sheet.protection.password, "the lock marks a boundary; it is not a lock"

    async def test_the_vendors_close_is_not_in_it(self, written: dict[str, Any]) -> None:
        """A computed figure may leave; the price series may not (ADR 0030, ADR 0034)."""
        book = self._book(written)
        for sheet in book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.value not in (Decimal(_CLOSE), float(_CLOSE), _CLOSE), (
                        f"{sheet.title}!{cell.coordinate} carries the close"
                    )

    async def test_the_market_capitalisation_that_weights_the_rate_is(
        self, written: dict[str, Any]
    ) -> None:
        sheet = self._book(written)["Model"]
        row = _row_labelled(sheet, "Equity value: market capitalisation")

        assert Decimal(str(sheet.cell(row=row, column=2).value)) == Decimal(_CAPITALISATION)

    async def test_the_sources_sheet_says_where_each_input_came_from(
        self, written: dict[str, Any]
    ) -> None:
        sheet = self._book(written)["Sources"]
        text = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)

        assert WHERE_THE_RECORD_ENDS in text
        assert written["report"].content_hash in text
        assert "Confirmed on" in text
        # The confirmer is recorded by address, and this file is mailed on.
        assert written["user"].email not in text
        assert "10-K for the period ending" in text
        # The scene's filer states its debt in parts, so the total is a sum, and says so.
        assert "Short-term debt plus long-term debt, as filed" in text
        assert "A filed line." not in text
        # F5 asks for each input's source document and the date it was retrieved.
        assert f"from {written['document'].url}" in text
        assert f"retrieved {written['document'].retrieved_at:%d %B %Y}" in text
        assert "The close itself is the price vendor's" in text

    async def test_the_grids_name_their_axes_as_the_model_sheet_names_the_inputs(
        self, written: dict[str, Any]
    ) -> None:
        """A reader holding a recorded grid finds the input it varies by the same words."""
        book = self._book(written)
        text = " ".join(
            str(cell.value) for row in book["Sensitivity"].iter_rows() for cell in row if cell.value
        )
        model = book["Model"]

        for label in ("Discount rate (WACC)", "Perpetual growth after the forecast"):
            assert label in text
            _row_labelled(model, label)
        assert "wacc" not in text
        assert "terminal_growth" not in text

    async def test_the_grids_are_the_reports_recorded_values(self, written: dict[str, Any]) -> None:
        sheet = self._book(written)["Sensitivity"]
        values = [
            cell.value
            for row in sheet.iter_rows(min_row=4)
            for cell in row
            if isinstance(cell.value, Decimal | float | int)
        ]

        assert values, "the recorded grids are carried"
        assert not any(
            isinstance(cell.value, str) and cell.value.startswith("=")
            for row in sheet.iter_rows()
            for cell in row
        ), "a recorded grid is values, not a second model"


class TestTheWorkbookSaysWhereThePlatformWouldRefuse:
    async def test_a_growth_rate_at_the_discount_rate_is_flagged_not_hidden(
        self, written: dict[str, Any], tmp_path: Path
    ) -> None:
        """The sheet computes whatever it is given; `gordon_terminal_value` would refuse."""
        book = _recalculated(
            _set_inputs(written["bytes"], "Perpetual growth after the forecast", Decimal("0.5")),
            tmp_path,
        )
        sheet = book["Model"]
        row = _row_labelled(sheet, "Would the platform value this?")

        assert str(sheet.cell(row=row, column=2).value).startswith("No: growth")
        assert sheet.cell(row=row, column=3).value == "Yes"

    async def test_the_model_as_the_report_struck_it_passes(
        self, written: dict[str, Any], tmp_path: Path
    ) -> None:
        sheet = _recalculated(written["bytes"], tmp_path)["Model"]
        row = _row_labelled(sheet, "Would the platform value this?")

        assert [sheet.cell(row=row, column=column).value for column in (2, 3)] == ["Yes", "Yes"]


def _multiple(key: str, label: str, value: str | None, *, absent: str = "") -> MultipleResult:
    return MultipleResult(
        key=key,
        label=label,
        quantity=None if value is None else Quantity.of(Decimal(value), DIMENSIONLESS),
        basis=MultipleBasis.LAST_FISCAL_YEAR,
        period_end=date(2024, 12, 31),
        absent_because=absent,
    )


def _table(*, publishable: bool) -> CompsTable:
    return CompsTable(
        subject=PeerRow(
            identifier="CTSO",
            name="Contoso Corporation",
            period_end=date(2024, 12, 31),
            multiples=(
                _multiple("ev_ebitda", "EV/EBITDA", "14.2"),
                _multiple("pe", "P/E", None, absent="earnings were negative"),
            ),
        ),
        peers=(
            PeerRow(
                identifier="FABR",
                name="Fabrikam Inc",
                period_end=date(2024, 12, 31),
                multiples=(_multiple("ev_ebitda", "EV/EBITDA", "11.8"),),
                rationale="The same end market, at a similar scale.",
            ),
        ),
        excluded=(
            PeerExclusion(
                identifier="NWND",
                name="Northwind Traders",
                reason="its prices were not fetched",
                rationale="A distributor to the same customers.",
            ),
        ),
        basis=MultipleBasis.LAST_FISCAL_YEAR,
        as_of=date(2025, 6, 30),
        peer_set_confirmed=True,
        licence_note="Computed figures may be published.",
        derived_figures_publishable=publishable,
    )


class TestTheComparisonGoesAsTheShareableReportPrintsIt:
    """A workbook is mailed on, so its comparables are the shareable audience's (ADR 0034)."""

    def test_a_publishable_table_carries_each_multiple_or_why_it_has_none(self) -> None:
        found = recorded_comps(_table(publishable=True).for_audience(Audience.SHAREABLE))

        assert found is not None
        assert found.columns == ("EV/EBITDA", "P/E")
        assert [line.name for line in found.lines] == [
            "Contoso Corporation (the company researched)",
            "Fabrikam Inc",
        ]
        assert found.lines[0].figures == (
            Decimal("14.2"),
            "Not meaningful: earnings were negative",
        )
        assert found.lines[1].figures == (Decimal("11.8"), "Not computed")
        assert ("Fabrikam Inc", "The same end market, at a similar scale.") in found.chosen
        assert found.left_out == ("Northwind Traders: its prices were not fetched",)

    def test_a_withheld_table_carries_the_companies_and_no_figure(self) -> None:
        found = recorded_comps(_table(publishable=False).for_audience(Audience.SHAREABLE))

        assert found is not None
        assert found.lines == ()
        assert "not reproduced here" in found.note
        assert [name for name, _ in found.chosen] == ["Fabrikam Inc", "Northwind Traders"]

    def test_no_comparison_is_none(self) -> None:
        assert recorded_comps(None) is None

    async def test_the_sheet_prints_the_figures_as_values(self, written: dict[str, Any]) -> None:
        comps = recorded_comps(_table(publishable=True).for_audience(Audience.SHAREABLE))
        model = replace(written["model"], comps=comps)

        sheet = load_workbook(BytesIO(build_workbook(model)))["Comparables"]
        figures = [
            cell.value
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, float | int | Decimal)
        ]

        assert sorted(Decimal(str(value)) for value in figures) == [
            Decimal("11.8"),
            Decimal("14.2"),
        ]

    async def test_a_withheld_sheet_prints_no_figure(self, written: dict[str, Any]) -> None:
        comps = recorded_comps(_table(publishable=False).for_audience(Audience.SHAREABLE))
        model = replace(written["model"], comps=comps)

        sheet = load_workbook(BytesIO(build_workbook(model)))["Comparables"]
        text = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)

        assert not any(
            isinstance(cell.value, float | int | Decimal)
            for row in sheet.iter_rows()
            for cell in row
        )
        assert "Northwind Traders" in text
        assert "A distributor to the same customers." in text

    async def test_a_run_that_compared_nobody_says_so(self, written: dict[str, Any]) -> None:
        sheet = load_workbook(BytesIO(written["bytes"]))["Comparables"]

        assert sheet.cell(row=3, column=1).value == "This run compared the company with nobody."


class TestTheSameRecordGivesTheSameBytes:
    """The PDF's rule: the archive's digest can be re-checked by writing the file again."""

    async def test_writing_it_again_later_gives_identical_bytes(
        self, written: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """An hour on by the clock the archive reads, and not a byte different."""
        later = time.time() + 3600
        monkeypatch.setattr(
            zipfile, "time", SimpleNamespace(time=lambda: later, localtime=time.localtime)
        )

        assert build_workbook(written["model"]) == written["bytes"]

    async def test_every_date_in_it_is_the_approvals(self, written: dict[str, Any]) -> None:
        """A wall clock anywhere in the file would make two writes differ."""
        approved = written["model"].approved_at.astimezone(UTC).replace(tzinfo=None)

        with ZipFile(BytesIO(written["bytes"])) as archive:
            stamps = {member.date_time for member in archive.infolist()}
        properties = load_workbook(BytesIO(written["bytes"])).properties

        assert stamps == {approved.timetuple()[:6]}
        assert properties.created == approved
        assert properties.modified == approved


class TestTheRenderStepArchivesIt:
    async def test_approval_archives_the_workbook_the_download_will_serve(
        self, written: dict[str, Any], tmp_path: Path
    ) -> None:
        """Archived when the report is approved, so a download is the file that was frozen
        rather than one written again from rows that may since have moved."""
        session: AsyncSession = written["session"]
        store = LocalArtefactStore(tmp_path / "artefacts", max_bytes=50_000_000)
        render = JobStep(
            job_id=written["job"].id,
            step_key="render",
            sequence=99,
            status=JobStatus.RUNNING,
            attempt=0,
            idempotency_key=f"{written['job'].id}:render",
            input_hash="b" * 64,
        )
        session.add(render)
        await session.flush()
        context = StepContext(
            session=session, job=written["job"], step=render, services={"store": store}
        )

        note = await _archive_workbook(context, store, report=written["report"])

        report: Report = written["report"]
        assert report.workbook_artefact_id is not None
        artefact = await session.get(Artefact, report.workbook_artefact_id)
        assert artefact is not None
        assert artefact.media_type == WORKBOOK_MEDIA_TYPE
        assert note == {"sha256": artefact.sha256}
        assert await store.read(artefact.sha256) == written["bytes"]


class TestWhenThereIsNoWorkbook:
    async def test_a_run_that_did_not_value_has_none(self, db_session: AsyncSession) -> None:
        scene = await valued_run(db_session, value=False)
        report = await _approved(db_session, scene)

        found = await workbook_model(db_session, job=scene["job"], report=report)

        assert isinstance(found, NoWorkbook)
        assert "no discounted cash flow" in found.reason

    async def test_an_unapproved_report_has_none(self, db_session: AsyncSession) -> None:
        scene = await valued_run(db_session)
        report = await _approved(db_session, scene)
        report.approved_at = None

        found = await workbook_model(db_session, job=scene["job"], report=report)

        assert isinstance(found, NoWorkbook)
        assert "only when a report is approved" in found.reason

    async def test_a_run_whose_filings_have_moved_has_none(self, db_session: AsyncSession) -> None:
        """Struck again, the model no longer gives the report's figures: a workbook written
        now would be a different model with the report's name on it."""
        scene = await valued_run(db_session)
        report = await _approved(db_session, scene)
        await seed_years(
            scene,
            {
                datetime(2025, 12, 31, tzinfo=UTC).date(): a_year(
                    revenue="1600", operating_income="420", **SHARES
                )
            },
        )

        found = await workbook_model(db_session, job=scene["job"], report=report)

        assert isinstance(found, NoWorkbook)
        assert "no longer gives the report's own figures" in found.reason
