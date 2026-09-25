"""The model workbook: the report's discounted cash flow with its formulas written out.

ADR 0134, F5. The report carries its valuation as figures whose footnotes lead to their
formulas, which is the right form for a record and the wrong one for a reader who wants to
change the growth rate and watch the answer move. This is the form for that reader: the same
model, as a spreadsheet whose every computed cell is a formula over cells they can change.

**The formulas are `aer.calc`'s own, written in the sheet's syntax, and nothing else.** Each
block below names the traced function it mirrors. A test recomputes the finished workbook in a
real spreadsheet engine and holds every figure to the ledger's, so a formula here that drifted
from its function would fail rather than ship a second, quietly different model.

**The modelling convention, because the audience knows it:** inputs blue and unlocked,
formulas black and locked, the sheet protected with no password. The lock marks the boundary
between what the platform recorded and what the reader is changing; it does not stop anybody
crossing it.

Pure: a :class:`ModelWorkbook` in, the file's bytes out. What goes into one — and what may
not, the vendor's price among it — is :mod:`aer.services.workbook`'s to decide.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Final
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.writer.excel import ExcelWriter

from aer.calc.dcf import MAX_EXIT_MULTIPLE

__all__ = [
    "WORKBOOK_MEDIA_TYPE",
    "CompsLine",
    "CostOfDebt",
    "Input",
    "ModelWorkbook",
    "RecordedComps",
    "RecordedGrid",
    "build_workbook",
]

WORKBOOK_MEDIA_TYPE: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# The sentence ADR 0134 puts on the sources sheet and the model sheet both: where the record
# ends. A workbook that did not say it would let an edited copy pass for the platform's own.
WHERE_THE_RECORD_ENDS: Final = (
    "The moment any figure here is changed, this is your model and no longer the platform's "
    "record: the report and its footnotes are the record."
)

_BLUE: Final = Font(color="0000FF")
_BLACK: Final = Font(color="000000")
# The report's own figures, held as values beside the model: neither an input nor a formula,
# so neither colour the convention gives those.
_RECORDED: Final = Font(color="7F7F7F")
_BOLD: Final = Font(bold=True)
_NOTE: Final = Font(italic=True, color="595959")
_UNLOCKED: Final = Protection(locked=False)

# Number formats by what a figure is. Amounts stay in the filing's own units: rescaling
# them into millions would make the sheet's arithmetic and the filing disagree by a factor
# nobody wrote down.
_FORMATS: Final[dict[str, str]] = {
    "amount": "#,##0;-#,##0",
    "per_share": "#,##0.00;-#,##0.00",
    "rate": "0.00%",
    "multiple": '0.00"\N{MULTIPLICATION SIGN}"',
    "count": "#,##0",
    "coefficient": "0.000",
    "factor": "0.0000",
}


@dataclass(frozen=True, slots=True)
class Input:
    """One blue cell: what it is, its value, and where the value came from."""

    label: str
    value: Decimal
    # "amount", "per_share", "rate", "multiple", "count" or "coefficient".
    style: str
    # One line for the sources sheet: the filing, the confirmation or the calculation.
    source: str
    kind: str


@dataclass(frozen=True, slots=True)
class CostOfDebt:
    """How the pre-tax cost of debt is reached — derived from filed lines, or confirmed.

    Derived where the filings carry an interest expense, as :func:`aer.calc.wacc.cost_of_debt`
    derives it: interest over the average of the prior and latest debt, or over the latest
    alone where there is no prior year. Confirmed where they do not. Neither where the company
    carries no debt, and then the discount rate is the cost of equity alone.
    """

    interest_expense: Input | None = None
    prior_debt: Input | None = None
    confirmed_rate: Input | None = None


@dataclass(frozen=True, slots=True)
class RecordedGrid:
    """One sensitivity grid as the report recorded it: values, not formulas."""

    title: str
    row_label: str
    column_label: str
    row_values: tuple[Decimal, ...]
    column_values: tuple[Decimal, ...]
    # Row-major, one tuple per row value.
    cells: tuple[tuple[Decimal, ...], ...]
    row_style: str
    column_style: str


@dataclass(frozen=True, slots=True)
class CompsLine:
    """One company in the comparison: each multiple, or the words for why it has none."""

    name: str
    figures: tuple[Decimal | str, ...]


@dataclass(frozen=True, slots=True)
class RecordedComps:
    """The comparison as the shareable report prints it: values, never the prices behind them.

    ``lines`` is empty where the licence withholds the multiples, and then ``note`` says so;
    the companies a person chose, and why, are listed either way (ADR 0034's amendment).
    """

    heading: str
    columns: tuple[str, ...]
    lines: tuple[CompsLine, ...]
    chosen: tuple[tuple[str, str], ...]
    left_out: tuple[str, ...]
    note: str


@dataclass(frozen=True, slots=True)
class ModelWorkbook:
    """Everything the workbook shows, assembled and checked by :mod:`aer.services.workbook`."""

    company: str
    ticker: str
    currency: str
    as_of: date
    run_id: str
    report_id: str
    content_hash: str
    code_version: str
    approved_at: datetime

    base_revenue: Input
    opening_working_capital: Input
    net_debt: Input
    shares_outstanding: Input
    tax_rate: Input
    terminal_growth: Input
    exit_multiple: Input
    risk_free_rate: Input
    beta: Input
    equity_risk_premium: Input
    equity_value: Input
    debt_value: Input
    cost_of_debt: CostOfDebt

    revenue_growth: tuple[Input, ...]
    ebit_margin: tuple[Input, ...]
    depreciation_intensity: tuple[Input, ...]
    capex_intensity: tuple[Input, ...]
    working_capital_intensity: tuple[Input, ...]

    # The report's own figures, by name, for the column beside the live ones.
    recorded: Mapping[str, Decimal]
    grids: tuple[RecordedGrid, ...] = ()
    # None where the run compared the company with nobody.
    comps: RecordedComps | None = None

    @property
    def years(self) -> int:
        return len(self.revenue_growth)


def build_workbook(model: ModelWorkbook) -> bytes:
    """The workbook's bytes: the model, the recorded grids, and the sources."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Model"
    _model_sheet(sheet, model)
    _grid_sheet(book.create_sheet("Sensitivity"), model.grids)
    _comps_sheet(book.create_sheet("Comparables"), model.comps)
    _sources_sheet(book.create_sheet("Sources"), model)

    # Stamped with the approval rather than the moment of writing, so the file says when the
    # record it carries was agreed rather than when a renderer happened to run.
    approved = model.approved_at.astimezone(UTC).replace(tzinfo=None)
    book.properties.creator = "Tracework Invest"
    book.properties.title = f"{model.company} — discounted cash flow"
    book.properties.created = approved
    book.properties.modified = approved

    # The writer is driven directly because `Workbook.save` stamps the modified date with the
    # wall clock on the way out. It closes the archive when it has written every part.
    buffer = BytesIO()
    ExcelWriter(book, ZipFile(buffer, "w", ZIP_DEFLATED, allowZip64=True)).save()
    return _restamped(buffer.getvalue(), approved)


def _restamped(data: bytes, when: datetime) -> bytes:
    """The same archive with every member dated ``when`` rather than the moment it was saved.

    The PDF's rule (:mod:`aer.render.pdf`): the same rows produce the same bytes, so the
    artefact's digest can be re-checked by writing the file again. openpyxl dates each member
    of the zip with the wall clock, which alone made two writes a second apart differ.
    """
    stamp = when.timetuple()[:6]
    out = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(out, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            dated = ZipInfo(member.filename, date_time=stamp)
            dated.compress_type = ZIP_DEFLATED
            dated.external_attr = member.external_attr
            target.writestr(dated, source.read(member))
    return out.getvalue()


# -- The model sheet ---------------------------------------------------------------------------


class _Cursor:
    """Writes rows down a sheet and remembers where each named figure landed."""

    def __init__(self, sheet: Worksheet) -> None:
        self.sheet = sheet
        self.row = 1
        self.at: dict[str, str] = {}

    def note(self, text: str, *, bold: bool = False) -> None:
        cell = self.sheet.cell(row=self.row, column=1, value=text)
        cell.font = _BOLD if bold else _NOTE
        self.row += 1

    def gap(self) -> None:
        self.row += 1

    def input(self, key: str, item: Input) -> str:
        """A blue, unlocked input in column B, labelled in A and sourced in C."""
        self.sheet.cell(row=self.row, column=1, value=item.label)
        cell = self.sheet.cell(row=self.row, column=2, value=item.value)
        cell.font = _BLUE
        cell.protection = _UNLOCKED
        cell.number_format = _FORMATS[item.style]
        self.sheet.cell(row=self.row, column=3, value=item.kind).font = _NOTE
        return self._remember(key)

    def formula(self, key: str, label: str, formula: str, *, style: str) -> str:
        """A black, locked formula in column B."""
        self.sheet.cell(row=self.row, column=1, value=label)
        cell = self.sheet.cell(row=self.row, column=2, value=formula)
        cell.font = _BLACK
        cell.number_format = _FORMATS[style]
        return self._remember(key)

    def _remember(self, key: str) -> str:
        address = f"$B${self.row}"
        self.at[key] = address
        self.row += 1
        return address


def _model_sheet(sheet: Worksheet, model: ModelWorkbook) -> None:
    at = _Cursor(sheet)
    title = f"{model.company} ({model.ticker})" if model.ticker else model.company
    at.note(f"Discounted cash flow — {title}", bold=True)
    at.note(
        f"As of {model.as_of:%d %B %Y}. Amounts in {model.currency} as filed. Blue cells are "
        "inputs, black cells are formulas over them, and grey cells are the report's own "
        "recorded figures."
    )
    at.note(WHERE_THE_RECORD_ENDS)
    at.gap()

    at.note("Inputs", bold=True)
    for key, item in (
        ("base_revenue", model.base_revenue),
        ("opening_working_capital", model.opening_working_capital),
        ("net_debt", model.net_debt),
        ("shares", model.shares_outstanding),
        ("tax_rate", model.tax_rate),
        ("terminal_growth", model.terminal_growth),
        ("exit_multiple", model.exit_multiple),
        ("risk_free", model.risk_free_rate),
        ("beta", model.beta),
        ("premium", model.equity_risk_premium),
        ("equity_value", model.equity_value),
        ("debt_value", model.debt_value),
    ):
        at.input(key, item)
    debt = model.cost_of_debt
    if debt.interest_expense is not None:
        at.input("interest", debt.interest_expense)
        if debt.prior_debt is not None:
            at.input("prior_debt", debt.prior_debt)
    elif debt.confirmed_rate is not None:
        at.input("confirmed_debt_rate", debt.confirmed_rate)
    at.gap()

    wacc = _cost_of_capital(at, model)
    at.gap()
    table = _forecast(sheet, at, model, wacc=wacc)
    at.gap()
    _terminal(sheet, at, model, table=table, wacc=wacc)

    sheet.column_dimensions["A"].width = 52
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 22
    for column in range(4, 4 + model.years):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    sheet.protection.sheet = True


def _cost_of_capital(at: _Cursor, model: ModelWorkbook) -> str:
    """`aer.calc.wacc`: the cost of equity, the cost of debt, the weights and the WACC."""
    at.note("Cost of capital", bold=True)
    r = at.at
    equity_cost = at.formula(
        "cost_of_equity",
        "Cost of equity",
        f"={r['risk_free']}+{r['beta']}*{r['premium']}",
        style="rate",
    )
    debt = model.cost_of_debt
    if debt.interest_expense is None and debt.confirmed_rate is None:
        # `wacc_all_equity`: no debt, so no cost of debt and no weights to strike.
        return at.formula("wacc", "Discount rate (WACC)", f"={equity_cost}", style="rate")

    if debt.interest_expense is not None:
        base = (
            f"(({r['prior_debt']}+{r['debt_value']})/2)"
            if debt.prior_debt is not None
            else r["debt_value"]
        )
        pre_tax = at.formula(
            "cost_of_debt", "Cost of debt, before tax", f"={r['interest']}/{base}", style="rate"
        )
    else:
        pre_tax = at.formula(
            "cost_of_debt",
            "Cost of debt, before tax",
            f"={r['confirmed_debt_rate']}",
            style="rate",
        )
    after_tax = at.formula(
        "after_tax_cost_of_debt",
        "Cost of debt, after tax",
        f"={pre_tax}*(1-{r['tax_rate']})",
        style="rate",
    )
    total = f"({r['equity_value']}+{r['debt_value']})"
    equity_weight = at.formula(
        "equity_weight", "Equity weight", f"={r['equity_value']}/{total}", style="rate"
    )
    debt_weight = at.formula(
        "debt_weight", "Debt weight", f"={r['debt_value']}/{total}", style="rate"
    )
    return at.formula(
        "wacc",
        "Discount rate (WACC)",
        f"={equity_cost}*{equity_weight}+{after_tax}*{debt_weight}",
        style="rate",
    )


@dataclass(frozen=True, slots=True)
class _Table:
    """Where the forecast's rows landed: the row number of each line, and its columns."""

    rows: dict[str, int]
    first: int  # the column of year one
    last: int  # the column of the final year

    def cell(self, line: str, column: int) -> str:
        return f"{get_column_letter(column)}{self.rows[line]}"

    def final(self, line: str) -> str:
        return f"${get_column_letter(self.last)}${self.rows[line]}"

    def span(self, line: str) -> str:
        return f"{self.cell(line, self.first)}:{self.cell(line, self.last)}"


def _forecast(sheet: Worksheet, at: _Cursor, model: ModelWorkbook, *, wacc: str) -> _Table:
    """`aer.calc.dcf.project`: each line of each year, in the order the function strikes it.

    Column C is year zero — the latest filed revenue and the opening working capital — and
    columns D onwards are the forecast years, so every year's formula reads the column to
    its left for "the year before".
    """
    at.note("Forecast", bold=True)
    header = at.row
    sheet.cell(row=header, column=1, value="Year").font = _BOLD
    sheet.cell(row=header, column=3, value=0).font = _BOLD
    first, last = 4, 3 + model.years
    for column in range(first, last + 1):
        sheet.cell(row=header, column=column, value=column - 3).font = _BOLD
    at.row += 1

    rows: dict[str, int] = {}
    r = at.at

    def line(key: str, label: str) -> int:
        rows[key] = at.row
        sheet.cell(row=at.row, column=1, value=label)
        at.row += 1
        return rows[key]

    def inputs(key: str, label: str, items: Sequence[Input]) -> None:
        row = line(key, label)
        for column, item in zip(range(first, last + 1), items, strict=True):
            cell = sheet.cell(row=row, column=column, value=item.value)
            cell.font = _BLUE
            cell.protection = _UNLOCKED
            cell.number_format = _FORMATS[item.style]

    def formulas(key: str, label: str, text: str, *, style: str, zero: str | None = None) -> None:
        """One line of formulas, a year per column.

        ``text`` is a template. ``{this}`` and ``{before}`` are the year's column and the
        one to its left, ``{row}`` is this line's own row, ``{header}`` the row holding the
        year numbers, and every line already written is a placeholder for its row — so
        ``={this}{revenue}*{this}{margin}`` reads as the function it mirrors.
        """
        row = line(key, label)
        if zero is not None:
            cell = sheet.cell(row=row, column=3, value=zero)
            cell.font = _BLACK
            cell.number_format = _FORMATS[style]
        for column in range(first, last + 1):
            written = text.format(
                this=get_column_letter(column),
                before=get_column_letter(column - 1),
                row=row,
                header=header,
                **rows,
            )
            cell = sheet.cell(row=row, column=column, value=written)
            cell.font = _BLACK
            cell.number_format = _FORMATS[style]

    tax = r["tax_rate"]
    inputs("growth", "Revenue growth", model.revenue_growth)
    # Each line below mirrors the traced function of the same idea in `aer.calc.dcf`:
    # projected revenue from the year before, then EBIT, NOPAT and the three intensities.
    formulas(
        "revenue",
        "Revenue",
        "={before}{row}*(1+{this}{growth})",
        style="amount",
        zero=f"={r['base_revenue']}",
    )
    inputs("margin", "EBIT margin", model.ebit_margin)
    formulas("ebit", "EBIT", "={this}{revenue}*{this}{margin}", style="amount")
    formulas("nopat", "NOPAT (EBIT after tax)", "={this}{ebit}*(1-" + tax + ")", style="amount")
    inputs(
        "depreciation_intensity",
        "Depreciation, as a share of revenue",
        model.depreciation_intensity,
    )
    formulas(
        "depreciation",
        "Depreciation",
        "={this}{revenue}*{this}{depreciation_intensity}",
        style="amount",
    )
    inputs("capex_intensity", "Capital expenditure, as a share of revenue", model.capex_intensity)
    formulas(
        "capex", "Capital expenditure", "={this}{revenue}*{this}{capex_intensity}", style="amount"
    )
    inputs(
        "working_capital_intensity",
        "Working capital, as a share of revenue",
        model.working_capital_intensity,
    )
    formulas(
        "working_capital",
        "Working capital",
        "={this}{revenue}*{this}{working_capital_intensity}",
        style="amount",
        zero=f"={r['opening_working_capital']}",
    )
    formulas(
        "working_capital_change",
        "Change in working capital",
        "={this}{working_capital}-{before}{working_capital}",
        style="amount",
    )
    formulas("ebitda", "EBITDA", "={this}{ebit}+{this}{depreciation}", style="amount")
    formulas(
        "fcff",
        "Free cash flow to the firm",
        "={this}{nopat}+{this}{depreciation}-{this}{capex}-{this}{working_capital_change}",
        style="amount",
    )
    # Discounted at the end of each year, as `discount_factor` and `present_value` do.
    formulas("factor", "Discount factor", "=1/(1+" + wacc + ")^{this}${header}", style="factor")
    formulas("present_value", "Present value", "={this}{fcff}*{this}{factor}", style="amount")
    return _Table(rows=rows, first=first, last=last)


def _terminal(
    sheet: Worksheet, at: _Cursor, model: ModelWorkbook, *, table: _Table, wacc: str
) -> None:
    """`aer.calc.dcf`'s two terminal methods, carried to a value per share side by side."""
    r = at.at
    at.note("Terminal value and value per share", bold=True)
    header = at.row
    for column, text in ((2, "Perpetuity growth"), (3, "Exit multiple")):
        sheet.cell(row=header, column=column, value=text).font = _BOLD
    at.row += 1

    fcff, ebitda = table.final("fcff"), table.final("ebitda")
    factor = table.final("factor")
    present = f"SUM({table.span('present_value')})"

    def pair(label: str, gordon: str, exit_multiple: str, *, style: str) -> int:
        row = at.row
        sheet.cell(row=row, column=1, value=label)
        for column, text in ((2, gordon), (3, exit_multiple)):
            if not text:
                continue
            cell = sheet.cell(row=row, column=column, value=text)
            cell.font = _BLACK
            cell.number_format = _FORMATS[style]
        at.row += 1
        return row

    tv = pair(
        "Terminal value",
        f"={fcff}*(1+{r['terminal_growth']})/({wacc}-{r['terminal_growth']})",
        f"={ebitda}*{r['exit_multiple']}",
        style="amount",
    )
    discounted = pair(
        "Discounted terminal value", f"=B{tv}*{factor}", f"=C{tv}*{factor}", style="amount"
    )
    enterprise = pair(
        "Enterprise value",
        f"={present}+B{discounted}",
        f"={present}+C{discounted}",
        style="amount",
    )
    pair(
        "Terminal value's share of it",
        f"=B{discounted}/B{enterprise}",
        f"=C{discounted}/C{enterprise}",
        style="rate",
    )
    equity = pair(
        "Equity value",
        f"=B{enterprise}-{r['net_debt']}",
        f"=C{enterprise}-{r['net_debt']}",
        style="amount",
    )
    per_share = pair(
        "Value per share",
        f"=B{equity}/{r['shares']}",
        f"=C{equity}/{r['shares']}",
        style="per_share",
    )
    # The report's own answer beneath the live one, as values: the row a reader compares
    # against once they have changed something, and the one that says the sheet reproduces
    # the report before they have.
    recorded = at.row
    sheet.cell(row=recorded, column=1, value="Value per share, as the report recorded it")
    for column, key in ((2, "gordon_per_share"), (3, "exit_multiple_per_share")):
        figure = model.recorded.get(key)
        if figure is None:
            continue
        cell = sheet.cell(row=recorded, column=column, value=figure)
        cell.font = _RECORDED
        cell.number_format = _FORMATS["per_share"]
    at.row += 1
    # The sheet computes whatever it is given, where `aer.calc.dcf` refuses a terminal value
    # that is arithmetic rather than a valuation. Its two terminal refusals are written out, so
    # a reader who crosses one is told rather than handed a number the report could not hold.
    growth, multiple = r["terminal_growth"], r["exit_multiple"]
    pair(
        "Would the platform value this?",
        (
            f'=IF({growth}>={wacc},"No: growth is not below the discount rate",'
            f'IF({fcff}<=0,"No: the final year of cash flow is not positive","Yes"))'
        ),
        (
            f'=IF({ebitda}<=0,"No: the final year of EBITDA is not positive",'
            f"IF(OR({multiple}<=0,{multiple}>{MAX_EXIT_MULTIPLE}),"
            f'"No: the multiple is outside 0 to {MAX_EXIT_MULTIPLE}","Yes"))'
        ),
        style="amount",
    )
    pair(
        "Exit multiple the perpetuity method implies",
        f"=B{tv}/{ebitda}",
        "",
        style="multiple",
    )
    pair(
        "Perpetual growth the exit multiple implies",
        "",
        f"=(C{tv}*{wacc}-{fcff})/(C{tv}+{fcff})",
        style="rate",
    )
    gap = at.row
    sheet.cell(row=gap, column=1, value="How far apart the two finish, as a share of the lower")
    cell = sheet.cell(
        row=gap,
        column=2,
        value=(
            f"=(MAX(B{per_share},C{per_share})-MIN(B{per_share},C{per_share}))"
            f"/MIN(B{per_share},C{per_share})"
        ),
    )
    cell.font = _BLACK
    cell.number_format = _FORMATS["rate"]
    at.row += 1
    for row in range(header, at.row):
        sheet.cell(row=row, column=1).alignment = Alignment(wrap_text=True)


# -- The recorded grids ------------------------------------------------------------------------


def _grid_sheet(sheet: Worksheet, grids: Sequence[RecordedGrid]) -> None:
    """The report's sensitivity grids, as recorded: values, because they are the record."""
    sheet.cell(row=1, column=1, value="Sensitivity, as the report recorded it").font = _BOLD
    sheet.cell(
        row=2,
        column=1,
        value=(
            "These cells are the report's own, one recorded calculation each. They do not move "
            "when the model sheet changes; change an input there to see its answer instead."
        ),
    ).font = _NOTE
    row = 4
    if not grids:
        sheet.cell(row=row, column=1, value="This report recorded no sensitivity grid.")
        return
    for grid in grids:
        sheet.cell(row=row, column=1, value=grid.title).font = _BOLD
        row += 1
        sheet.cell(
            row=row, column=1, value=f"Down: {grid.row_label}. Across: {grid.column_label}."
        ).font = _NOTE
        for offset, value in enumerate(grid.column_values, start=2):
            cell = sheet.cell(row=row, column=offset, value=value)
            cell.font = _BOLD
            cell.number_format = _FORMATS[grid.column_style]
        row += 1
        for value, cells in zip(grid.row_values, grid.cells, strict=True):
            head = sheet.cell(row=row, column=1, value=value)
            head.font = _BOLD
            head.number_format = _FORMATS[grid.row_style]
            for offset, figure in enumerate(cells, start=2):
                cell = sheet.cell(row=row, column=offset, value=figure)
                cell.font = _RECORDED
                cell.number_format = _FORMATS["per_share"]
            row += 1
        row += 1
    sheet.column_dimensions["A"].width = 40
    sheet.protection.sheet = True


# -- The comparison ------------------------------------------------------------------------------


def _comps_sheet(sheet: Worksheet, comps: RecordedComps | None) -> None:
    """The report's comparison with other companies, as values: it is the record, not a model."""
    sheet.cell(row=1, column=1, value="Comparables, as the report shows them").font = _BOLD
    row = 3
    if comps is None:
        sheet.cell(row=row, column=1, value="This run compared the company with nobody.")
        sheet.protection.sheet = True
        return
    sheet.cell(row=2, column=1, value=comps.note).font = _NOTE
    sheet.cell(row=row, column=1, value=comps.heading).font = _BOLD
    row += 1
    if comps.lines:
        for column, text in enumerate(("Company", *comps.columns), start=1):
            sheet.cell(row=row, column=column, value=text).font = _BOLD
        row += 1
        for line in comps.lines:
            sheet.cell(row=row, column=1, value=line.name)
            for column, figure in enumerate(line.figures, start=2):
                cell = sheet.cell(row=row, column=column, value=figure)
                cell.font = _RECORDED
                if isinstance(figure, Decimal):
                    cell.number_format = _FORMATS["multiple"]
                else:
                    cell.alignment = Alignment(wrap_text=True)
            row += 1
    row += 1
    for heading, entries in (
        ("Why each company was chosen", [f"{name}: {why}" for name, why in comps.chosen]),
        ("Not in the table", list(comps.left_out)),
    ):
        if not entries:
            continue
        sheet.cell(row=row, column=1, value=heading).font = _BOLD
        row += 1
        for entry in entries:
            sheet.cell(row=row, column=1, value=entry).alignment = Alignment(wrap_text=True)
            row += 1
        row += 1
    sheet.column_dimensions["A"].width = 60
    for column in range(2, 2 + len(comps.columns)):
        sheet.column_dimensions[get_column_letter(column)].width = 16
    sheet.protection.sheet = True


# -- The sources ---------------------------------------------------------------------------------


def _sources_sheet(sheet: Worksheet, model: ModelWorkbook) -> None:
    """Every input with where it came from, and the record this workbook was taken from."""
    sheet.cell(row=1, column=1, value="Where every input came from").font = _BOLD
    sheet.cell(row=2, column=1, value=WHERE_THE_RECORD_ENDS).font = _NOTE
    stamp = (
        ("Report", model.report_id),
        ("Report content hash (SHA-256)", model.content_hash),
        ("Run", model.run_id),
        ("Code version", model.code_version),
        ("Approved", f"{model.approved_at.astimezone(UTC):%d %B %Y, %H:%M} UTC"),
    )
    row = 4
    for label, value in stamp:
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=value)
        row += 1
    row += 1
    for column, text in enumerate(("Input", "Value", "Kind", "Where it came from"), start=1):
        sheet.cell(row=row, column=column, value=text).font = _BOLD
    row += 1

    debt = model.cost_of_debt
    listed = (
        model.base_revenue,
        model.opening_working_capital,
        model.net_debt,
        model.shares_outstanding,
        model.tax_rate,
        model.terminal_growth,
        model.exit_multiple,
        model.risk_free_rate,
        model.beta,
        model.equity_risk_premium,
        model.equity_value,
        model.debt_value,
        *(
            item
            for item in (debt.interest_expense, debt.prior_debt, debt.confirmed_rate)
            if item is not None
        ),
        *model.revenue_growth,
        *model.ebit_margin,
        *model.depreciation_intensity,
        *model.capex_intensity,
        *model.working_capital_intensity,
    )
    for item in listed:
        sheet.cell(row=row, column=1, value=item.label)
        figure = sheet.cell(row=row, column=2, value=item.value)
        figure.number_format = _FORMATS[item.style]
        sheet.cell(row=row, column=3, value=item.kind)
        sheet.cell(row=row, column=4, value=item.source).alignment = Alignment(wrap_text=True)
        row += 1
    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 22
    sheet.column_dimensions["D"].width = 90
    sheet.protection.sheet = True
